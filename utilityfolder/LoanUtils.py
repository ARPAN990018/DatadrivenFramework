import openpyxl


def rownumber(path,sheetname):
    workbook = openpyxl.load_workbook(path)
    sheet = workbook[sheetname]
    return sheet.max_row

def columnnumber(path,sheetname):
    workbook = openpyxl.load_workbook(path)
    sheet = workbook[sheetname]
    return sheet.max_column

def readdata(path,sheetname,rownum,columnnum):
    workbook = openpyxl.load_workbook(path)
    sheet = workbook[sheetname]
    return sheet.cell(rownum,columnnum).value

def writedata(path,sheetname,rownum,columnnum,data):
    workbook = openpyxl.load_workbook(path)
    sheet = workbook[sheetname]
    sheet.cell(rownum, columnnum).value = data
    workbook.save(path)

